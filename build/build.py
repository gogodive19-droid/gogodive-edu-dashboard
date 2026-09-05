# -*- coding: utf-8 -*-
"""data/raw/*.xlsx → 정규화 → 중복 처리 → 집계 → dist/data.json + dist/index.html

사용:  python build/build.py
규칙(사용자와 확정):
  - 인원 = '출석' 컬럼. 0명 수업은 셀로 보존하되 UI에서 횟수에 넣지 않는다.
  - 같은 (수업일시, 강사명, 수업명) 이 여러 줄이면 출석이 가장 적은 1건만 남긴다.
  - '프리다이빙 이론 [LIVE]' 는 제외.
  - 수업명 규칙에 없는 값이 하나라도 있으면 목록을 출력하고 실패한다.
"""
import glob
import json
import os
import sys
from collections import Counter
from datetime import datetime

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from mapping import BIG_ORDER, UnmappedClassName, parse_class_name  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_DIR = os.path.join(ROOT, "data", "raw")
DIST_DIR = os.path.join(ROOT, "dist")
TEMPLATE = os.path.join(ROOT, "src", "template.html")
SHEET = "그룹수업"
HEADER = ("수업일시", "강사명", "수업명", "출석 + 결석", "출석", "결석")
SIZE_CAP = 11   # 11명 이상은 UI에서 한 칸으로 묶지만 데이터에는 실제 인원을 남긴다


def _count(cell):
    """'8명' → 8, 빈 값 → 0."""
    if cell is None:
        return 0
    return int(str(cell).replace("명", "").strip() or 0)


def load_raw_rows(paths):
    """엑셀 파일들 → [(수업일시 'YYYY-MM-DD HH:MM', 강사, 수업명, 출석)]"""
    rows = []
    for path in paths:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = tuple(str(c).strip() if c is not None else "" for c in next(it)[:6])
        if header != HEADER:
            raise SystemExit("헤더가 예상과 다릅니다: %s\n  %s" % (path, header))
        for r in it:
            if r[0] is None or str(r[0]).strip() == "":
                continue
            dt = str(r[0]).strip()[:16]
            rows.append((dt, str(r[1]).strip(), str(r[2]).strip(), _count(r[4])))
        wb.close()
    return rows


def dedupe(rows):
    """같은 (수업일시, 강사, 수업명) → 출석 최솟값 1건. 반환 (rows, 제거된 줄 수)"""
    best = {}
    for dt, inst, name, n in rows:
        key = (dt, inst, name)
        if key not in best or n < best[key]:
            best[key] = n
    return [(k[0], k[1], k[2], n) for k, n in best.items()], len(rows) - len(best)


def normalize(rows):
    """행 → 레코드 dict. 반환 (records, excluded_count, unmapped Counter)"""
    records, excluded, unmapped = [], 0, Counter()
    for dt, inst, name, n in rows:
        try:
            p = parse_class_name(name)
        except UnmappedClassName:
            unmapped[name] += 1
            continue
        if p is None:
            excluded += 1
            continue
        records.append({"m": dt[:7], "d": dt[:10], "i": inst,
                        "b": p.big, "s": p.small, "v": p.venue, "n": n})
    return records, excluded, unmapped


def aggregate(records, source_files):
    """레코드 → 사전 배열 + 셀 [월, 강사, 세부수업, 강습장, 인원, 횟수]"""
    months = sorted({r["m"] for r in records})
    # 강사: 진행 횟수(0명 제외) 많은 순
    inst_cnt = Counter(r["i"] for r in records if r["n"] > 0)
    instructors = sorted({r["i"] for r in records}, key=lambda x: (-inst_cnt[x], x))
    # 세부 수업: 대분류 순서 → 횟수 많은 순
    cls_cnt = Counter((r["b"], r["s"]) for r in records if r["n"] > 0)
    cls_all = {(r["b"], r["s"]) for r in records}
    classes = sorted(cls_all, key=lambda bs: (BIG_ORDER.index(bs[0]), -cls_cnt[bs], bs[1]))
    bigs = [b for b in BIG_ORDER if any(c[0] == b for c in classes)]
    ven_cnt = Counter(r["v"] for r in records if r["n"] > 0)
    venues = sorted({r["v"] for r in records}, key=lambda x: (-ven_cnt[x], x))

    mi = {m: k for k, m in enumerate(months)}
    ii = {x: k for k, x in enumerate(instructors)}
    ci = {c: k for k, c in enumerate(classes)}
    vi = {v: k for k, v in enumerate(venues)}

    cells = Counter()
    for r in records:
        cells[(mi[r["m"]], ii[r["i"]], ci[(r["b"], r["s"])], vi[r["v"]], r["n"])] += 1
    cell_list = [list(k) + [c] for k, c in sorted(cells.items())]

    return {
        "meta": {
            "builtAt": datetime.now().strftime("%Y-%m-%d %H:%M"),
            "firstMonth": months[0], "lastMonth": months[-1],
            "sourceFiles": [os.path.basename(p) for p in source_files],
            "records": len(records),
            "sizeCap": SIZE_CAP,
        },
        "months": months,
        "instructors": instructors,
        "classes": [{"s": s, "b": b} for b, s in classes],
        "bigs": bigs,
        "venues": venues,
        "cells": cell_list,
    }


def render(data):
    os.makedirs(DIST_DIR, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":"))
    with open(os.path.join(DIST_DIR, "data.json"), "w", encoding="utf-8") as f:
        f.write(payload)
    if not os.path.exists(TEMPLATE):
        print("template 없음 → dist/data.json 만 생성:", TEMPLATE)
        return None
    with open(TEMPLATE, encoding="utf-8") as f:
        html = f.read()
    marker = "/*__DATA__*/"
    if marker not in html:
        raise SystemExit("template.html 에 %s 자리표시자가 없습니다" % marker)
    html = html.replace(marker, payload.replace("</script", "<\\/script"), 1)
    out = os.path.join(DIST_DIR, "index.html")
    with open(out, "w", encoding="utf-8") as f:
        f.write(html)
    return out


def main():
    paths = sorted(glob.glob(os.path.join(RAW_DIR, "*.xlsx")))
    if not paths:
        raise SystemExit("data/raw/ 에 엑셀 파일이 없습니다")
    raw = load_raw_rows(paths)
    rows, removed = dedupe(raw)
    records, excluded, unmapped = normalize(rows)
    print("원본 행 %d  →  중복 제거 %d  →  제외(이론 LIVE 등) %d  →  레코드 %d"
          % (len(raw), removed, excluded, len(records)))
    if unmapped:
        print("\n!! 규칙에 없는 수업명 %d종 — build/mapping.py 에 규칙을 추가하세요:" % len(unmapped))
        for name, c in unmapped.most_common():
            print("   %5d  %s" % (c, name))
        raise SystemExit(1)
    data = aggregate(records, paths)
    out = render(data)
    print("기간 %s ~ %s | 강사 %d | 세부수업 %d | 강습장 %d | 셀 %d"
          % (data["meta"]["firstMonth"], data["meta"]["lastMonth"], len(data["instructors"]),
             len(data["classes"]), len(data["venues"]), len(data["cells"])))
    if out:
        print("생성:", out, "(%.0f KB)" % (os.path.getsize(out) / 1024))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
